#!/usr/bin/env python3
"""
tools/update_srs_docx.py — write real benchmark numbers into the SRS.

Kept separate from tools/benchmark_srs_table.py (which measures) so
re-measuring never requires touching this file and vice versa.

Edits kkwieer/SRS/RAPtwin_SRS_v1.2.docx in place:
  1. Finds the §3.2.4 benchmark table by content (first data row says
     "DDQN-HTRCS"), not by a hard-coded index — fails loudly if the structure
     has drifted since this was written, rather than silently editing the
     wrong table.
  2. Replaces the 4 real-strategy rows with measured numbers from
     reports/srs_benchmark_table.json (written by benchmark_srs_table.py).
  3. Deletes the DDQN-HTRCS row — nothing in this codebase implements it.
  4. Rewrites the intro paragraph (found by substring match on "edge devices")
     to describe the real fabric instead of "M=35 edge devices, N=4 edge
     servers".
  5. Inserts a caption paragraph after the table with the exact reproduction
     recipe (commit, seed, commands) so the numbers are independently
     checkable.

Also flips FR-28, FR-29, FR-31's status cell from "Planned" to "Implemented"
(FR-30 is left untouched — still Planned, see future.md).

Usage:
    python -m tools.update_srs_docx
    python -m tools.update_srs_docx --no-pdf   # skip the soffice conversion
"""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
from collections.abc import Sequence
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

DOCX_PATH = ROOT / "kkwieer" / "SRS" / "RAPtwin_SRS_v1.2.docx"
TABLE_DATA_PATH = ROOT / "reports" / "srs_benchmark_table.json"
FLIP_TO_IMPLEMENTED = ("FR-28", "FR-29", "FR-31")


def set_cell_text(cell: Any, text: str) -> None:
    """Replace a table cell's text, preserving run formatting where possible."""
    paragraphs = cell.paragraphs
    if len(paragraphs) == 1 and len(paragraphs[0].runs) == 1:
        paragraphs[0].runs[0].text = text
        return
    # Fallback: clear and write a plain run. Only hit for cells this script
    # doesn't expect to touch in the current document (format drift guard).
    for p in paragraphs[1:]:
        p._p.getparent().remove(p._p)
    paragraphs[0].clear()
    paragraphs[0].add_run(text)


def find_benchmark_table(doc: Any) -> Any:
    matches = [t for t in doc.tables if any("DDQN" in r.cells[0].text for r in t.rows)]
    if len(matches) != 1:
        raise SystemExit(
            f"expected exactly one table containing a 'DDQN' row, found {len(matches)} — "
            "the SRS structure has drifted since this script was written; fix the search "
            "condition before editing anything."
        )
    return matches[0]


def update_table(table: Any, rows: list) -> None:
    by_strategy = {r["strategy"]: r for r in rows}
    data_rows = list(table.rows[1:])  # skip header
    ddqn_row = next((r for r in data_rows if "DDQN" in r.cells[0].text), None)
    if ddqn_row is None:
        raise SystemExit(
            "could not re-locate the DDQN-HTRCS row inside the matched table"
        )

    for row in data_rows:
        label = row.cells[0].text.strip()
        if label not in by_strategy:
            continue  # the DDQN row — handled below
        r = by_strategy[label]
        p95 = "—" if r["p95_jct_s"] is None else f"{r['p95_jct_s']}"
        set_cell_text(row.cells[1], f"{r['mean_jct_s']}")
        set_cell_text(row.cells[2], p95)
        set_cell_text(row.cells[3], f"{r['energy_j']}")
        set_cell_text(row.cells[4], f"{r['sla_hit_rate_pct']}%")

    ddqn_row._tr.getparent().remove(ddqn_row._tr)


def update_intro_paragraph(doc: Any, measured: dict[str, Any]) -> None:
    target = next((p for p in doc.paragraphs if "edge devices" in p.text), None)
    if target is None:
        raise SystemExit(
            "could not find the §3.2.4 intro paragraph (searched for 'edge devices')"
        )
    new_text = (
        f"The following benchmark results (reference fabric: nodes/*.yaml, "
        f"job catalogue jobs/jobs_10.yaml limited to the first {measured['job_limit']} jobs, "
        f"under the '{measured['scenario']}' chaos scenario at deadline_scale="
        f"{measured['deadline_scale']}) illustrate the resilience gap between baseline "
        f"heuristics and the resilient/RL-based planners:"
    )
    if len(target.runs) == 1:
        target.runs[0].text = new_text
    else:
        for r in target.runs[1:]:
            r.text = ""
        target.runs[0].text = new_text


def insert_caption_after_table(doc: Any, table: Any, measured: dict[str, Any]) -> None:
    caption = (
        f"Measured at commit {measured['commit']}, seed {measured['seed']} "
        f"({measured['duration_s']}s). {measured['note']} Reproduce with: "
        f"python -m tools.benchmark_srs_table --seed {measured['seed']}."
    )
    new_para = doc.add_paragraph(caption)
    new_para.style = (
        doc.styles["Caption"]
        if "Caption" in [s.name for s in doc.styles]
        else new_para.style
    )
    table._tbl.addnext(new_para._p)


def update_fr_status(doc: Any) -> int:
    """Flip FR-28/29/31's status cell to Implemented. FR-30 stays Planned."""
    flipped = 0
    for table in doc.tables:
        header = [c.text.strip() for c in table.rows[0].cells]
        if header[:2] != ["ID", "Requirement"]:
            continue
        for row in table.rows[1:]:
            fr_id = row.cells[0].text.strip()
            if fr_id in FLIP_TO_IMPLEMENTED and row.cells[-1].text.strip() == "Planned":
                set_cell_text(row.cells[-1], "Implemented")
                flipped += 1
    return flipped


def convert_to_pdf(docx_path: Path) -> Path | None:
    out_dir = docx_path.parent
    result = subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out_dir),
            str(docx_path),
        ],
        capture_output=True,
        text=True,
    )
    if result.returncode != 0:
        print(
            f"PDF conversion failed:\n{result.stdout}\n{result.stderr}", file=sys.stderr
        )
        return None
    pdf_path = out_dir / (docx_path.stem + ".pdf")
    return pdf_path if pdf_path.exists() else None


def build_argparser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(
        description="Write measured benchmark numbers into the SRS docx."
    )
    ap.add_argument("--data", default=str(TABLE_DATA_PATH))
    ap.add_argument("--docx", default=str(DOCX_PATH))
    ap.add_argument(
        "--no-pdf", action="store_true", help="skip regenerating the PDF export"
    )
    return ap


def main(argv: Sequence[str] | None = None) -> int:
    import docx  # imported here so --help works without the dependency installed

    args = build_argparser().parse_args(argv)
    data_path = Path(args.data)
    docx_path = Path(args.docx)
    if not data_path.exists():
        raise SystemExit(
            f"{data_path} not found — run `python -m tools.benchmark_srs_table` first"
        )
    if not docx_path.exists():
        raise SystemExit(f"{docx_path} not found")

    measured = json.loads(data_path.read_text(encoding="utf-8"))

    doc = docx.Document(str(docx_path))
    table = find_benchmark_table(doc)
    update_table(table, measured["rows"])
    update_intro_paragraph(doc, measured)
    insert_caption_after_table(doc, table, measured)
    flipped = update_fr_status(doc)
    doc.save(str(docx_path))
    print(
        f"Updated {docx_path.name}: table rewritten, DDQN-HTRCS row removed, "
        f"{flipped} FR status cell(s) flipped to Implemented."
    )

    if not args.no_pdf:
        pdf_path = convert_to_pdf(docx_path)
        if pdf_path:
            print(f"Regenerated {pdf_path}")
        else:
            print(
                "PDF regeneration failed — see stderr above; docx was still updated.",
                file=sys.stderr,
            )
            return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
